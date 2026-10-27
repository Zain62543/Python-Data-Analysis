from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, PP_PARAGRAPH_ALIGNMENT
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

def set_slide_background(slide, rgb_hex):
    """Set slide background color"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(rgb_hex)

def add_code_slide(pres, title_text, code_text):
    """Add a slide with code formatting"""
    slide = pres.slides.add_slide(pres.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    
    # Create textbox for code
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(4)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    
    # Format as code
    p = tf.add_paragraph()
    p.text = code_text
    p.font.name = 'Courier New'
    p.font.size = Pt(14)
    p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT
    
    return slide

# Create presentation
prs = Presentation()

# ========== Slide 1: Title Slide ==========
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Object-Oriented Programming in Python"
subtitle.text = "Complete Guide with Examples\nPresented by [Your Name]"
set_slide_background(slide, "2a5885")  # Dark blue

# ========== Slide 2: Agenda ==========
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Agenda"
content.text = (
    "1. Introduction to OOP Principles\n"
    "2. Classes & Objects in Python\n"
    "3. Inheritance & Polymorphism\n"
    "4. Encapsulation & Abstraction\n"
    "5. Special Methods & Operator Overloading\n"
    "6. Practical OOP Examples\n"
    "7. Best Practices & Common Pitfalls"
)

# ========== Slide 3: What is OOP? ==========
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "What is Object-Oriented Programming?"
content.text = (
    "• Programming paradigm based on 'objects'\n"
    "• Objects contain:\n"
    "  - Data (attributes/properties)\n"
    "  - Behavior (methods/functions)\n"
    "• Key Principles (PIE):\n"
    "  - Polymorphism\n"
    "  - Inheritance\n"
    "  - Encapsulation\n"
    "  - Abstraction (sometimes considered separate)"
)

# ========== Slide 4: Class Definition ==========
add_code_slide(
    prs,
    "Defining a Class in Python",
    "class Dog:\n"
    "    species = 'Canis familiaris'  # Class attribute\n\n"
    "    def __init__(self, name, age):\n"
    "        self.name = name  # Instance attribute\n"
    "        self.age = age\n\n"
    "    def bark(self):\n"
    "        return f'{self.name} says woof!'\n\n"
    "# Creating an instance\n"
    "my_dog = Dog('Rex', 4)\n"
    "print(my_dog.bark())"
)

# ========== Slide 5: __init__ Method ==========
add_code_slide(
    prs,
    "The __init__ Method",
    "class Car:\n"
    "    def __init__(self, make, model, year):\n"
    "        self.make = make\n"
    "        self.model = model\n"
    "        self.year = year\n"
    "        self.odometer = 0  # Default value\n\n"
    "    def get_description(self):\n"
    "        return f'{self.year} {self.make} {self.model}'"
)

# ========== Slide 6: Instance vs Class Attributes ==========
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Instance vs Class Attributes"
content.text = (
    "Instance Attributes:\n"
    "• Unique to each object\n"
    "• Defined in __init__\n"
    "• Accessed via self.attribute\n\n"
    "Class Attributes:\n"
    "• Shared by all instances\n"
    "• Defined outside __init__\n"
    "• Accessed via ClassName.attribute"
)

# ========== Slide 7: Inheritance ==========
add_code_slide(
    prs,
    "Inheritance Basics",
    "class Animal:  # Parent class\n"
    "    def __init__(self, name):\n"
    "        self.name = name\n\n"
    "    def speak(self):\n"
    "        raise NotImplementedError\n\n"
    "class Dog(Animal):  # Child class\n"
    "    def speak(self):\n"
    "        return f'{self.name} says woof!'\n\n"
    "class Cat(Animal):\n"
    "    def speak(self):\n"
    "        return f'{self.name} says meow!'"
)

# ========== Slide 8: super() ==========
add_code_slide(
    prs,
    "Using super()",
    "class Person:\n"
    "    def __init__(self, name):\n"
    "        self.name = name\n\n"
    "class Employee(Person):\n"
    "    def __init__(self, name, id):\n"
    "        super().__init__(name)  # Calls parent __init__\n"
    "        self.employee_id = id"
)

# ========== Slide 9: Polymorphism ==========
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Polymorphism"
content.text = (
    "Same interface, different implementations:\n"
    "1. Method Overriding (inheritance)\n"
    "2. Duck Typing (Python's approach)\n\n"
    "Example:\n"
    "def animal_sound(animal):\n"
    "    print(animal.speak())\n\n"
    "Works with any object that has speak() method!"
)

# ========== Continue with remaining slides... ==========

# Save presentation
prs.save("python_oop_presentation.pptx")
print("Presentation created successfully!")