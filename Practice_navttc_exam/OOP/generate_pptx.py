from pptx import Presentation
from pptx.util import Inches

# Create a presentation object
prs = Presentation()

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Introduction to Object-Oriented Programming in Python"
subtitle.text = "Presented by: [Your Name]\nDate: August 16, 2025\nOverview: Learn the fundamentals of OOP in Python"

# Add more slides (example for Slide 2)
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Understanding Object-Oriented Programming"
content = slide.placeholders[1]
content.text = "A programming paradigm based on objects and classes\nFocuses on data encapsulation, inheritance, and polymorphism\nBenefits: Modularity, reusability, and scalability"

# Add remaining slides similarly (omitted for brevity)
# Save the presentation
prs.save("OOP_Python_Presentation.pptx")